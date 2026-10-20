import csv
import io
import os
import re
from datetime import date, datetime, timedelta

from sqlalchemy import event

from library_app.extensions import db
from library_app.models import AuditLog, Book, BookCopy, BugIncident, Complaint, Feedback, Loan, Notification, PasswordResetToken, SystemSetting, User, utcnow
from library_app.services import refresh_loan_notifications, time_series_monthly
from library_app.time_utils import library_today

from .conftest import login


def test_personal_ui_calendar_help_and_about_render(client, records, app):
    login(client, "member@test.edu", "MemberPassword!2026")
    for path, phrase in [("/calendar", b"Due-date calendar"), ("/help", b"frequently asked questions"), ("/about", b"Campus Library")]:
        response = client.get(path)
        assert response.status_code == 200
        assert phrase in response.data
    shell = client.get("/dashboard")
    assert b"Display preferences" in shell.data
    assert b"Search books, authors, ISBN" in shell.data
    assert b"High contrast" in shell.data
    css_version = format(os.stat(os.path.join(app.static_folder, "css", "app.css")).st_mtime_ns, "x")
    js_version = format(os.stat(os.path.join(app.static_folder, "js", "app.js")).st_mtime_ns, "x")
    assert f"/static/css/app.css?v={css_version}".encode() in shell.data
    assert f"/static/js/app.js?v={js_version}".encode() in shell.data


def test_catalog_matches_card_based_stitch_layout_with_real_inventory(client, records):
    login(client, "member@test.edu", "MemberPassword!2026")
    response = client.get("/library/books")
    assert response.status_code == 200
    assert b"Explore the collection" in response.data
    assert b"data-catalog-view=\"grid\"" in response.data
    assert b"data-catalog-view=\"table\"" in response.data
    assert b"Test Book" in response.data
    assert b"1 / 1 available" in response.data
    assert b"1,428" not in response.data


def test_staff_dashboard_renders_visual_board_from_real_records(client, records):
    login(client, "admin@test.edu", "AdminPassword!2026")
    response = client.get("/dashboard")
    assert response.status_code == 200
    for phrase in [b"YOUR CAMPUS COLLECTION", b"Borrowing rhythm", b"Copies at a glance",
                   b"Due this week", b"Find your next read", b"Test Book"]:
        assert phrase in response.data


def test_catalog_course_subject_semester_filters_and_local_book_photography(client, records):
    login(client, "member@test.edu", "MemberPassword!2026")
    response = client.get("/library/books", query_string={"course": records["course_id"],
                                                           "category": records["category_id"], "semester": 2})
    assert response.status_code == 200
    assert b"Test Book" in response.data
    assert b"Test BSc Course" in response.data
    assert b"SEM 2" in response.data
    assert b"images/books/" in response.data
    empty = client.get("/library/books", query_string={"course": records["course_id"], "semester": 4})
    assert b"No books found on this shelf" in empty.data


def test_system_settings_persist_and_record_audit(client, records, app):
    login(client, "admin@test.edu", "AdminPassword!2026")
    response = client.post("/system-settings", data={"default_loan_days": "21", "fine_per_day": "2.50"}, follow_redirects=True)
    assert response.status_code == 200
    assert b"System settings saved" in response.data
    with app.app_context():
        assert db.session.get(SystemSetting, "DEFAULT_LOAN_DAYS").value == "21"
        assert db.session.get(SystemSetting, "FINE_PER_DAY").value == "2.50"
        assert AuditLog.query.filter_by(action="update_system_settings").count() == 1
    invalid = client.post("/system-settings", data={"default_loan_days": "0", "fine_per_day": "-1"}, follow_redirects=True)
    assert b"must be between 1 and 365" in invalid.data


def test_password_reset_is_expiring_hashed_and_one_use(client, records, app, monkeypatch):
    from library_app import auth

    captured = {}
    monkeypatch.setattr(auth, "send_password_reset_email", lambda user, token: captured.update(email=user.email, token=token) or True)
    response = client.post("/password-reset", data={"email": "member@test.edu"}, follow_redirects=True)
    assert response.status_code == 200
    assert b"If an active account matches" in response.data
    with app.app_context():
        token_row = PasswordResetToken.query.one()
        assert token_row.token_hash != captured["token"]
        assert token_row.expires_at > token_row.created_at
    response = client.post(f"/password-reset/{captured['token']}", data={
        "password": "NewMemberPassword!2026", "confirm_password": "NewMemberPassword!2026",
    }, follow_redirects=True)
    assert b"Your password has been reset" in response.data
    with app.app_context():
        assert db.session.get(User, records["member_id"]).check_password("NewMemberPassword!2026")
        assert db.session.get(User, records["member_id"]).auth_version == 1
        assert PasswordResetToken.query.one().used_at is not None
    replay = client.post(f"/password-reset/{captured['token']}", data={
        "password": "AttackerPassword!2026", "confirm_password": "AttackerPassword!2026",
    }, follow_redirects=True)
    assert b"invalid or has expired" in replay.data
    with app.app_context():
        assert db.session.get(User, records["member_id"]).check_password("NewMemberPassword!2026")


def test_admin_password_reset_increments_auth_version(app, records, monkeypatch):
    import getpass

    answers = iter(["UpdatedAdminPassword!2026", "UpdatedAdminPassword!2026"])
    monkeypatch.setattr(getpass, "getpass", lambda _prompt: next(answers))
    result = app.test_cli_runner().invoke(args=["set-password", "admin@test.edu"])
    assert result.exit_code == 0
    assert "Password updated." in result.output
    with app.app_context():
        admin = db.session.get(User, records["admin_id"])
        assert admin.auth_version == 1
        assert admin.check_password("UpdatedAdminPassword!2026")


def test_catalog_csv_import_requires_validated_preview(client, records, app):
    login(client, "admin@test.edu", "AdminPassword!2026")
    csv_text = ("isbn,title,author,category,publisher,publication_year,copies,course,semester\n"
                f"9780306406157,Imported Book,Test Author,Test Category,Test Publisher,2022,2,Test BSc Course,4\n")
    response = client.post("/library/books/import", data={
        "action": "upload", "file": (io.BytesIO(csv_text.encode()), "books.csv")
    }, content_type="multipart/form-data")
    assert response.status_code == 200
    assert b"1 valid row" in response.data
    batch_id = re.search(rb'name="batch_id" value="(\d+)"', response.data).group(1).decode()
    response = client.post("/library/books/import", data={"action": "import", "batch_id": batch_id}, follow_redirects=True)
    assert response.status_code == 200
    with app.app_context():
        book = Book.query.filter_by(isbn="9780306406157").one()
        assert book.title == "Imported Book"
        assert book.total_copies == 2
        assert book.course_id == records["course_id"] and book.semester == 4

    invalid_csv = ("isbn,title,author,category,publisher,publication_year,copies\n"
                   "9781234567890,Duplicate,Test Author,Test Category,Test Publisher,2020,1\n")
    response = client.post("/library/books/import", data={
        "action": "upload", "file": (io.BytesIO(invalid_csv.encode()), "duplicate.csv")
    }, content_type="multipart/form-data")
    assert b"Import paused" in response.data
    assert b"duplicate ISBN" in response.data


def test_catalog_import_caps_total_copies_before_any_books_are_added(client, records, app):
    login(client, "admin@test.edu", "AdminPassword!2026")

    def isbn13(sequence):
        body = f"978{sequence:09d}"
        check_digit = (10 - sum(int(digit) * (1 if index % 2 == 0 else 3)
                                for index, digit in enumerate(body)) % 10) % 10
        return body + str(check_digit)

    rows = [f"{isbn13(number)},Bulk title {number},Test Author,Test Category,Test Publisher,2024,500"
            for number in range(21)]
    csv_text = "isbn,title,author,category,publisher,publication_year,copies\n" + "\n".join(rows) + "\n"
    response = client.post("/library/books/import", data={
        "action": "upload", "file": (io.BytesIO(csv_text.encode()), "large-copy-batch.csv"),
    }, content_type="multipart/form-data")

    assert response.status_code == 200
    assert b"Import paused" in response.data
    assert b"10,000 copies in total" in response.data
    with app.app_context():
        assert Book.query.count() == 1


def test_incident_register_validates_status_and_is_role_limited(client, records, app):
    login(client, "quality@test.edu", "QualityPassword!2026")
    response = client.post("/incidents", data={"title": "Search results freeze", "description": "The catalog search did not return after submitting a category filter.", "severity": "high"}, follow_redirects=True)
    assert response.status_code == 200
    assert b"Incident added to the tracking log" in response.data
    with app.app_context():
        item = BugIncident.query.one()
        incident_id = item.id
        assert item.status == "open"
    client.post(f"/incidents/{incident_id}/status", data={"status": "resolved"}, follow_redirects=True)
    with app.app_context():
        item = db.session.get(BugIncident, incident_id)
        assert item.status == "resolved" and item.resolved_at is not None
    client.post("/logout")
    login(client, "member@test.edu", "MemberPassword!2026")
    assert client.get("/incidents").status_code == 403


def test_report_exports_csv_excel_pdf_and_role_permissions(client, records):
    login(client, "admin@test.edu", "AdminPassword!2026")
    period = {"start_date": date.today().isoformat(), "end_date": date.today().isoformat()}
    csv_response = client.get("/reports/export/inventory/csv", query_string=period)
    assert csv_response.status_code == 200 and csv_response.mimetype == "text/csv"
    assert b"Test Book" in csv_response.data
    xlsx_response = client.get("/reports/export/inventory/xlsx", query_string=period)
    assert xlsx_response.status_code == 200
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(xlsx_response.data), read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows[0][0] == "ISBN" and rows[1][1] == "Test Book"
    assert rows[0][3:6] == ("Course", "Subject", "Semester")
    assert rows[1][3] == "Test BSc Course" and rows[1][5] == 2
    pdf_response = client.get("/reports/export/inventory/pdf", query_string=period)
    assert pdf_response.status_code == 200 and pdf_response.data.startswith(b"%PDF")
    assert client.get("/reports/export/inventory/csv", query_string={"start_date": "2026-99-99"}).status_code == 400
    client.post("/logout")
    login(client, "quality@test.edu", "QualityPassword!2026")
    assert client.get("/reports/export/members/xlsx", query_string=period).status_code == 403
    assert client.get("/reports/export/quality/pdf", query_string=period).status_code == 200


def test_member_activity_export_uses_bounded_queries_for_many_members(client, records, app):
    login(client, "admin@test.edu", "AdminPassword!2026")
    with app.app_context():
        db.session.add_all([
            User(full_name=f"Extra Member {number:02}", email=f"extra{number}@test.edu",
                 role="member", password_hash="unused")
            for number in range(20)
        ])
        db.session.add(Loan(copy_id=records["copy_id"], member_id=records["member_id"],
                            issued_by_id=records["admin_id"], due_date=date.today()))
        db.session.add(Complaint(member_id=records["member_id"], category="other",
                                 subject="Report count check", description="Count this request."))
        db.session.commit()
        engine = db.engine

    statements = []

    def record_statement(_connection, _cursor, statement, _parameters, _context, _executemany):
        statements.append(statement)

    event.listen(engine, "before_cursor_execute", record_statement)
    try:
        response = client.get("/reports/members.csv")
    finally:
        event.remove(engine, "before_cursor_execute", record_statement)

    assert response.status_code == 200
    assert len([statement for statement in statements if statement.lstrip().upper().startswith("SELECT")]) <= 3
    rows = list(csv.reader(io.StringIO(response.data.decode("utf-8-sig"))))
    member_row = next(row for row in rows if row[0] == str(records["member_id"]))
    assert member_row[4:] == ["1", "1"]


def test_due_loan_notifications_use_bounded_queries_and_are_idempotent(app, records):
    with app.app_context():
        today = library_today()
        copies = [BookCopy(book_id=records["book_id"], barcode=f"REMINDER-{number}", status="loaned")
                  for number in range(5)]
        db.session.add_all(copies)
        db.session.flush()
        db.session.add_all([
            Loan(copy_id=copy.id, member_id=records["member_id"], issued_by_id=records["admin_id"],
                 due_date=today + timedelta(days=1))
            for copy in copies
        ])
        db.session.commit()
        engine = db.engine

    statements = []

    def record_statement(_connection, _cursor, statement, _parameters, _context, _executemany):
        statements.append(statement)

    event.listen(engine, "before_cursor_execute", record_statement)
    try:
        with app.app_context():
            refresh_loan_notifications(records["member_id"])
            refresh_loan_notifications(records["member_id"])
    finally:
        event.remove(engine, "before_cursor_execute", record_statement)

    selects = [statement for statement in statements if statement.lstrip().upper().startswith("SELECT")]
    assert len(selects) <= 4
    with app.app_context():
        reminders = Notification.query.filter_by(user_id=records["member_id"], kind="due").all()
        assert len(reminders) == 5
        assert len({reminder.reference_key for reminder in reminders}) == 5


def test_monthly_charts_aggregate_in_database_with_expected_counts_and_averages(app, records):
    with app.app_context():
        current_month = utcnow().replace(day=15, hour=12, minute=0, second=0, microsecond=0)
        previous_month = current_month - timedelta(days=25)
        since = current_month - timedelta(days=60)
        db.session.add_all([
            Feedback(member_id=records["member_id"], category="other", rating=5, comments="Previous month",
                     created_at=previous_month),
            Feedback(member_id=records["member_id"], category="other", rating=2, comments="Current month one",
                     created_at=current_month),
            Feedback(member_id=records["member_id"], category="other", rating=4, comments="Current month two",
                     created_at=current_month + timedelta(days=1)),
            Loan(copy_id=records["copy_id"], member_id=records["member_id"], issued_by_id=records["admin_id"],
                 due_date=library_today(), issued_at=previous_month, status="returned"),
            Loan(copy_id=records["copy_id"], member_id=records["member_id"], issued_by_id=records["admin_id"],
                 due_date=library_today(), issued_at=current_month, status="returned"),
            Loan(copy_id=records["copy_id"], member_id=records["member_id"], issued_by_id=records["admin_id"],
                 due_date=library_today(), issued_at=current_month + timedelta(days=1), status="returned"),
        ])
        db.session.commit()
        engine = db.engine

    statements = []

    def record_statement(_connection, _cursor, statement, _parameters, _context, _executemany):
        statements.append(statement)

    event.listen(engine, "before_cursor_execute", record_statement)
    try:
        with app.app_context():
            satisfaction = time_series_monthly(Feedback, Feedback.created_at, since,
                                               value_field=Feedback.rating,
                                               until=current_month + timedelta(days=10))
            borrowing = time_series_monthly(Loan, Loan.issued_at, since,
                                            until=current_month + timedelta(days=10))
    finally:
        event.remove(engine, "before_cursor_execute", record_statement)

    month_labels = [previous_month.strftime("%Y-%m"), current_month.strftime("%Y-%m")]
    assert satisfaction == {"labels": month_labels, "values": [5.0, 3.0]}
    assert borrowing == {"labels": month_labels, "values": [1, 2]}
    selects = [statement.upper() for statement in statements if statement.lstrip().upper().startswith("SELECT")]
    assert len(selects) == 2
    assert all("GROUP BY" in statement for statement in selects)
    assert any("AVG(" in statement for statement in selects)


def test_admin_pages_and_upload_permissions(client, records):
    login(client, "member@test.edu", "MemberPassword!2026")
    assert client.get("/system-settings").status_code == 403
    assert client.get("/audit-log").status_code == 403
    assert client.get("/library/books/import").status_code == 403
    client.post("/logout")
    login(client, "admin@test.edu", "AdminPassword!2026")
    for path in ["/system-settings", "/audit-log", "/incidents", "/library/books/import", "/reports/"]:
        assert client.get(path).status_code == 200, path


def test_expired_records_purge_and_audit(app, records):
    from library_app.models import CatalogImportBatch, PasswordResetToken, AuditLog, utcnow
    from library_app.services import expired_record_counts, purge_expired_records
    from datetime import timedelta

    with app.app_context():
        past = utcnow() - timedelta(hours=2)
        future = utcnow() + timedelta(hours=2)

        # Expired token and valid token
        expired_token = PasswordResetToken(user_id=1, token_hash="exp_hash", expires_at=past)
        valid_token = PasswordResetToken(user_id=1, token_hash="val_hash", expires_at=future)
        # Expired batch and valid batch
        expired_batch = CatalogImportBatch(owner_id=1, filename="old.csv", rows_json="[]", expires_at=past)
        valid_batch = CatalogImportBatch(owner_id=1, filename="new.csv", rows_json="[]", expires_at=future)

        db.session.add_all([expired_token, valid_token, expired_batch, valid_batch])
        db.session.commit()
        exp_tok_id = expired_token.id
        val_tok_id = valid_token.id
        exp_bat_id = expired_batch.id
        val_bat_id = valid_batch.id

        counts = expired_record_counts()
        assert counts["password_reset_tokens"] >= 1
        assert counts["catalog_import_previews"] >= 1

        db.session.expire_all()
        result = purge_expired_records(source="test_runner")
        assert result["password_reset_tokens"] >= 1
        assert result["catalog_import_previews"] >= 1

        # Check valid records remain
        assert db.session.get(PasswordResetToken, val_tok_id) is not None
        assert db.session.get(CatalogImportBatch, val_bat_id) is not None
        assert db.session.get(PasswordResetToken, exp_tok_id) is None
        assert db.session.get(CatalogImportBatch, exp_bat_id) is None

        audit = AuditLog.query.filter_by(action="purge_expired_records").first()
        assert audit is not None
        assert "source=test_runner" in audit.detail
