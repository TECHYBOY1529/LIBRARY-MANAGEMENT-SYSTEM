import csv
import io
from datetime import date, datetime, time, timedelta, timezone

from flask import Blueprint, Response, abort, current_app, render_template, request
from flask_login import current_user, login_required
from sqlalchemy import func
from sqlalchemy.orm import joinedload, selectinload
from .extensions import db
from .models import Book, BookCopy, Complaint, Feedback, Fine, Loan, User, utcnow
from .security import roles_required
from .services import kpi_summary, pareto_data, time_series_monthly

bp = Blueprint("reports", __name__)


@bp.get("/")
@roles_required("admin", "quality_manager")
def index():
    today = date.today()
    start_date, end_date = _date_range(default_start=today - timedelta(days=89), default_end=today)
    return render_template("reports.html", start_date=start_date, end_date=end_date)


def csv_response(filename, headers, rows):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(headers)
    def safe_cell(value):
        if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
            return "'" + value
        return value
    writer.writerows([[safe_cell(value) for value in row] for row in rows])
    return Response("\ufeff" + output.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f"attachment; filename={filename}"})


@bp.get("/inventory.csv")
@roles_required("admin", "quality_manager")
def inventory_csv():
    rows = Book.query.order_by(Book.title).all()
    return csv_response("inventory-report.csv", ["ISBN", "Title", "Author", "Category", "Publisher", "Year", "Total copies", "Available copies"],
                        [[b.isbn, b.title, b.author.name, b.category.name, b.publisher.name, b.publication_year, b.total_copies, b.available_copies] for b in rows])


@bp.get("/loans.csv")
@roles_required("admin")
def loans_csv():
    rows = Loan.query.order_by(Loan.issued_at.desc()).all()
    return csv_response("borrowing-report.csv", ["Loan ID", "Member", "ISBN", "Title", "Issued", "Due", "Returned", "Status", "Issue operation ms"],
                        [[r.id, r.member.full_name, r.copy.book.isbn, r.copy.book.title, r.issued_at.isoformat(), r.due_date.isoformat(), r.returned_at.isoformat() if r.returned_at else "", r.status, r.issue_seconds if r.issue_seconds is not None else ""] for r in rows])


@bp.get("/overdue.csv")
@roles_required("admin")
def overdue_csv():
    rows = Loan.query.filter(Loan.status == "active", Loan.due_date < date.today()).order_by(Loan.due_date).all()
    return csv_response("overdue-report.csv", ["Loan ID", "Member", "Email", "Title", "Due date", "Days overdue"],
                        [[r.id, r.member.full_name, r.member.email, r.copy.book.title, r.due_date.isoformat(), (date.today() - r.due_date).days] for r in rows])


@bp.get("/fines.csv")
@roles_required("admin")
def fines_csv():
    rows = Fine.query.join(Loan).order_by(Fine.assessed_at.desc()).all()
    return csv_response("fines-report.csv", ["Fine ID", "Member", "Title", "Amount", "Paid", "Assessed"],
                        [[r.id, r.loan.member.full_name, r.loan.copy.book.title, str(r.amount), r.paid, r.assessed_at.isoformat()] for r in rows])


@bp.get("/members.csv")
@roles_required("admin")
def members_csv():
    rows = User.query.filter_by(role="member").order_by(User.full_name).all()
    return csv_response("member-activity-report.csv", ["Member ID", "Name", "Email", "Active", "Loans", "Open complaints"],
                        [[r.id, r.full_name, r.email, r.active, Loan.query.filter_by(member_id=r.id).count(), Complaint.query.filter_by(member_id=r.id).filter(Complaint.status.in_(["open", "in_progress"])).count()] for r in rows])


@bp.get("/quality.csv")
@roles_required("admin", "quality_manager")
def quality_csv():
    summary = kpi_summary(90)
    labels = {
        "satisfaction": ("Average customer satisfaction", "/ 5"),
        "availability_when_requested_pct": ("Available when requested", "%"),
        "avg_issue_processing_ms": ("Average issue operation time", "ms"),
        "avg_complaint_resolution_days": ("Average complaint resolution time", "days"),
        "complaint_resolution_rate_pct": ("Complaint resolution rate", "%"),
        "overdue_loan_rate_pct": ("Overdue loan rate", "%"),
        "inventory_accuracy_pct": ("Inventory accuracy", "%"),
        "actions_on_time_pct": ("Actions completed on time", "%"),
        "borrowing_count": ("Borrowing activity", "loans"),
    }
    rows = [[label, summary[key] if summary[key] is not None else "Insufficient data", unit,
             "Point-in-time" if key == "overdue_loan_rate_pct" else "90 days",
             summary["data_sources"][key]] for key, (label, unit) in labels.items()]
    return csv_response("quality-kpi-report.csv", ["KPI", "Value", "Unit", "Reporting period", "Formula / data source"], rows)


@bp.get("/pareto.csv")
@roles_required("admin", "quality_manager")
def pareto_csv():
    data = pareto_data(90)
    rows = [[label, count, cumulative, data["priority_category"] == label.lower().replace(" ", "_")]
            for label, count, cumulative in zip(data["labels"], data["counts"], data["cumulative"])]
    return csv_response("problem-pareto-report.csv", ["Category", "Problem count", "Cumulative percent", "Most frequent"], rows)


@bp.get("/borrowing-trends.csv")
@roles_required("admin", "quality_manager")
def borrowing_trends_csv():
    since = utcnow() - timedelta(days=90)
    data = time_series_monthly(Loan, Loan.issued_at, since)
    return csv_response("borrowing-trends.csv", ["Month", "Loans issued"], zip(data["labels"], data["values"]))


REPORT_ACCESS = {
    "inventory": {"admin", "quality_manager"}, "loans": {"admin"}, "overdue": {"admin"},
    "fines": {"admin"}, "members": {"admin"}, "quality": {"admin", "quality_manager"},
    "pareto": {"admin", "quality_manager"}, "trends": {"admin", "quality_manager"},
}


def _date_range(default_start, default_end):
    def read(key, fallback):
        raw = request.args.get(key)
        if not raw:
            return fallback
        try:
            return date.fromisoformat(raw)
        except ValueError:
            abort(400, description="Enter dates in YYYY-MM-DD format.")
    start = read("start_date", default_start)
    end = read("end_date", default_end)
    if end < start or (end - start).days > 3660:
        abort(400, description="Choose a valid date range of no more than ten years.")
    return start, end


def _date_bounds(start, end):
    return (datetime.combine(start, time.min, tzinfo=timezone.utc),
            datetime.combine(end + timedelta(days=1), time.min, tzinfo=timezone.utc))


def _between(field, start, end):
    lower, upper = _date_bounds(start, end)
    return (field >= lower, field < upper)


def _report_dataset(kind, start, end):
    lower, upper = _date_bounds(start, end)
    if kind == "inventory":
        books = (Book.query.options(joinedload(Book.author), joinedload(Book.category), joinedload(Book.publisher),
                                    selectinload(Book.copies)).order_by(Book.title).limit(25001).all())
        rows = [[book.isbn, book.title, book.author.name, book.category.name, book.publisher.name,
                 book.publication_year, book.total_copies, book.available_copies] for book in books]
        return "Inventory report", ["ISBN", "Title", "Author", "Category", "Publisher", "Year", "Total copies", "Available copies"], rows
    if kind == "loans":
        records = (Loan.query.filter(*_between(Loan.issued_at, start, end))
                   .options(joinedload(Loan.member), joinedload(Loan.copy).joinedload(BookCopy.book))
                   .order_by(Loan.issued_at.desc()).limit(25001).all())
        rows = [[loan.id, loan.member.full_name, loan.copy.book.isbn, loan.copy.book.title,
                 loan.issued_at.isoformat(), loan.due_date.isoformat(), loan.returned_at.isoformat() if loan.returned_at else "",
                 loan.status, loan.issue_seconds if loan.issue_seconds is not None else ""] for loan in records]
        return "Borrowing and returns", ["Loan ID", "Member", "ISBN", "Title", "Issued", "Due", "Returned", "Status", "Issue operation ms"], rows
    if kind == "overdue":
        records = (Loan.query.filter(Loan.status == "active", Loan.due_date >= start, Loan.due_date <= end,
                                     Loan.due_date < date.today())
                   .options(joinedload(Loan.member), joinedload(Loan.copy).joinedload(BookCopy.book))
                   .order_by(Loan.due_date).limit(25001).all())
        rows = [[loan.id, loan.member.full_name, loan.member.email, loan.copy.book.title, loan.due_date.isoformat(),
                 (date.today() - loan.due_date).days] for loan in records]
        return "Overdue loans", ["Loan ID", "Member", "Email", "Title", "Due date", "Days overdue"], rows
    if kind == "fines":
        records = (Fine.query.join(Loan).filter(*_between(Fine.assessed_at, start, end))
                   .options(joinedload(Fine.loan).joinedload(Loan.member),
                            joinedload(Fine.loan).joinedload(Loan.copy).joinedload(BookCopy.book))
                   .order_by(Fine.assessed_at.desc()).limit(25001).all())
        rows = [[fine.id, fine.loan.member.full_name, fine.loan.copy.book.title, str(fine.amount), fine.paid,
                 fine.assessed_at.isoformat()] for fine in records]
        return "Fine register", ["Fine ID", "Member", "Title", "Amount", "Paid", "Assessed"], rows
    if kind == "members":
        members = User.query.filter_by(role="member").order_by(User.full_name).limit(25001).all()
        loan_counts = dict(db.session.query(Loan.member_id, func.count(Loan.id)).filter(*_between(Loan.issued_at, start, end)).group_by(Loan.member_id).all())
        complaint_counts = dict(db.session.query(Complaint.member_id, func.count(Complaint.id))
                                .filter(*_between(Complaint.created_at, start, end), Complaint.status.in_(["open", "in_progress"]))
                                .group_by(Complaint.member_id).all())
        rows = [[member.id, member.full_name, member.email, member.active, loan_counts.get(member.id, 0),
                 complaint_counts.get(member.id, 0)] for member in members]
        return "Member activity", ["Member ID", "Name", "Email", "Active", "Loans in period", "Open complaints in period"], rows
    if kind == "quality":
        summary = kpi_summary(start_date=start, end_date=end)
        labels = {
            "satisfaction": ("Average customer satisfaction", "/ 5"),
            "availability_when_requested_pct": ("Available when requested", "%"),
            "avg_issue_processing_ms": ("Average issue processing time", "ms"),
            "avg_complaint_resolution_days": ("Average complaint resolution time", "days"),
            "complaint_resolution_rate_pct": ("Complaint resolution rate", "%"),
            "overdue_loan_rate_pct": ("Overdue loan rate", "%"),
            "inventory_accuracy_pct": ("Inventory accuracy", "%"),
            "actions_on_time_pct": ("Actions completed on time", "%"),
            "borrowing_count": ("Borrowing activity", "loans"),
        }
        rows = [[label, summary[key] if summary[key] is not None else "Insufficient data", unit,
                 f"{start.isoformat()} to {end.isoformat()}", summary["data_sources"][key]]
                for key, (label, unit) in labels.items()]
        return "Quality KPI report", ["KPI", "Value", "Unit", "Reporting period", "Formula / data source"], rows
    if kind == "pareto":
        data = pareto_data(start_date=start, end_date=end)
        rows = [[label, count, cumulative, data["priority_category"] == label.lower().replace(" ", "_")]
                for label, count, cumulative in zip(data["labels"], data["counts"], data["cumulative"])]
        return "Problem-category Pareto", ["Category", "Problem count", "Cumulative percent", "Most frequent"], rows
    if kind == "trends":
        data = time_series_monthly(Loan, Loan.issued_at, lower, until=upper)
        return "Borrowing trends", ["Month", "Loans issued"], list(zip(data["labels"], data["values"]))
    abort(404)


@bp.get("/export/<kind>/<file_format>")
@login_required
def export_report(kind, file_format):
    if kind not in REPORT_ACCESS or current_user.role not in REPORT_ACCESS[kind]:
        abort(403)
    if file_format not in {"csv", "xlsx", "pdf"}:
        abort(404)
    today = date.today()
    start, end = _date_range(today - timedelta(days=89), today)
    title, headers, rows = _report_dataset(kind, start, end)
    if len(rows) > 25000:
        abort(413, description="Narrow the date range to export at most 25,000 rows.")
    stem = f"{kind}-report-{start.isoformat()}-to-{end.isoformat()}"
    if file_format == "csv":
        return csv_response(f"{stem}.csv", headers, rows)
    if file_format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet("Report")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{chr(64 + len(headers))}{len(rows) + 1}"
        header_style = {"font": Font(color="FFFFFF", bold=True), "fill": PatternFill("solid", fgColor="1B4332"),
                        "alignment": Alignment(vertical="center", wrap_text=True)}
        header_cells = []
        for value in headers:
            cell = WriteOnlyCell(worksheet, value=value)
            for key, style in header_style.items(): setattr(cell, key, style)
            header_cells.append(cell)
        worksheet.append(header_cells)
        for row in rows:
            safe_row = []
            for value in row:
                if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
                    value = "'" + value
                safe_row.append(value)
            worksheet.append(safe_row)
        output = io.BytesIO()
        workbook.save(output)
        return Response(output.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{stem}.xlsx"'})
    if len(rows) > 1500:
        abort(413, description="Narrow the date range to export at most 1,500 rows as PDF.")
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm,
                                 topMargin=12 * mm, bottomMargin=12 * mm, title=title)
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(title), styles["Title"]),
             Paragraph(f"Reporting period: {start.isoformat()} to {end.isoformat()} · Generated {utcnow().strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]),
             Spacer(1, 7 * mm)]
    table_rows = [[Paragraph(f"<b>{escape(str(value))}</b>", styles["Normal"]) for value in headers]]
    for row in rows:
        table_rows.append([Paragraph(escape(str(value if value is not None else "")), styles["BodyText"]) for value in row])
    width = landscape(A4)[0] - 24 * mm
    table = Table(table_rows, colWidths=[width / len(headers)] * len(headers), repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B4332")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#d8e2da")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f7f3")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    document.build(story)
    return Response(output.getvalue(), mimetype="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{stem}.pdf"'})
